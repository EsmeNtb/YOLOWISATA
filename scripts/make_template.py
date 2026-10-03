"""data/content.json -> templates/YoloWisata_Data_Template.xlsx
Run this when you want a fresh spreadsheet that matches what the website currently shows."""
import json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
sys.path.insert(0, str(Path(__file__).resolve().parent))
from columns import BIZ_COLS, CARD_COLS, Q_COLS, PHRASE_COLS

ROOT = Path(__file__).resolve().parent.parent
c = json.loads((ROOT / "data" / "content.json").read_text(encoding="utf-8"))
wb = Workbook()
HEAD = PatternFill("solid", start_color="FFC928")
F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)


def sheet(title, cols, data, widths):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for r in data:
        ws.append(r)
    for row in ws.iter_rows():
        for cell in row:
            cell.font = FB if cell.row == 1 else F
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row == 1:
                cell.fill = HEAD
    for i, wd in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wd
    ws.freeze_panes = "A2"


ws = wb.active
ws.title = "Read me first"
for line in [
    ["YoloWisata data template"],
    ["Fill in the four other sheets, save, then run:  python scripts/xlsx_to_data.py"],
    [""],
    ["Businesses", "One row per local business. id: short, lowercase, no spaces (noor, darto). sector: one of " + ", ".join(c["sectors"]) + ". map_x / map_y: 0 to 100 on the sketch map (0,0 is top left). chips: comma separated."],
    ["Postcards", "One row per visitor memory. business_id must match a row in Businesses. text: exactly as the visitor wrote it, in their language. english / indonesian: translations (leave empty when the text is already in that language). background: sun, field, sky or coral. local: 1 for a domestic visitor, empty otherwise."],
    ["Questions", "Things visitors asked before or during the visit. Same rules as Postcards."],
    ["Phrases", "The flip cards. pronunciation: spelled out for an English speaker."],
    [""],
    ["Consent", "Only add a real visitor's words with their permission. Use a nickname, never a full name. Mark invented rows as synthetic in your data sheet for the judges."],
    ["Counting", "A postcard is only counted in the owner's insights when it contains a keyword from data/content.json > themes. validate_data.py warns about rows that match nothing."],
]:
    ws.append(line)
for row in ws.iter_rows():
    for cell in row:
        cell.font = FB if (cell.row == 1 or cell.column == 1) else F
        cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 110

B = c["businesses"]
sheet("Businesses", BIZ_COLS, [[b["id"], b["name"], b["host"], b["sector"], b.get("e", ""), b.get("col", ""), b.get("place", ""),
       (b.get("pos") or ["", ""])[0], (b.get("pos") or ["", ""])[1], b.get("story", ""), b.get("L", {}).get("duration", ""), b.get("L", {}).get("price", ""),
       b.get("L", {}).get("availability", ""), b.get("L", {}).get("activities", ""), ", ".join(b.get("chips", []))] for b in B],
      [10, 26, 14, 10, 7, 10, 18, 7, 7, 60, 16, 22, 26, 40, 40])
sheet("Postcards", CARD_COLS, [[b["id"], k.get("n", ""), k.get("f", ""), k.get("l", ""), k["t"], k.get("en", ""), k.get("idt", ""), k.get("bg", ""), k.get("s", ""), 1 if k.get("loc") else ""] for b in B for k in b.get("cards", [])],
      [12, 12, 6, 12, 55, 55, 55, 12, 9, 6])
sheet("Questions", Q_COLS, [[b["id"], q.get("f", ""), q["t"], q.get("en", ""), q.get("idt", ""), 1 if q.get("loc") else ""] for b in B for q in b.get("qs", [])],
      [12, 6, 55, 55, 55, 6])
sheet("Phrases", PHRASE_COLS, [[p["id"], p["en"], p["say"], p["emoji"]] for p in c["phrases"]], [34, 34, 44, 7])
out = ROOT / "templates" / "YoloWisata_Data_Template.xlsx"
out.parent.mkdir(exist_ok=True)
wb.save(out)
print("Wrote", out)
