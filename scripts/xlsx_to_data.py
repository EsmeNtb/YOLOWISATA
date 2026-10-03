"""Data template spreadsheet -> data/content.json (+ refreshes the static copy and validates).

  python scripts/xlsx_to_data.py                      uses templates/YoloWisata_Data_Template.xlsx
  python scripts/xlsx_to_data.py path/to/file.xlsx

The spreadsheet replaces: businesses (with their postcards and questions) and phrases.
It keeps what is not in the spreadsheet: each existing business's Travel DNA values, "why it fits" lines,
seed bookings and seed message threads, plus sectors, swipe cards, themes and ideas.
"""
import json, subprocess, sys
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parent))
from columns import BIZ_COLS, CARD_COLS, Q_COLS, PHRASE_COLS

ROOT = Path(__file__).resolve().parent.parent
src = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "templates" / "YoloWisata_Data_Template.xlsx"
content_path = ROOT / "data" / "content.json"
c = json.loads(content_path.read_text(encoding="utf-8"))
wb = load_workbook(src, data_only=True)


def rows(sheet, cols):
    ws = wb[sheet]
    head = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    lost = [col for col in cols if col not in head]
    if lost:
        sys.exit("Sheet '%s' is missing column(s): %s" % (sheet, ", ".join(lost)))
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = {h: ("" if v is None else str(v).strip()) for h, v in zip(head, r) if h}
        if any(d.values()):
            yield d


def truthy(v):
    return str(v).strip().lower() in ("1", "1.0", "yes", "y", "true", "ya")


old = {b["id"]: b for b in c["businesses"]}
biz = []
for r in rows("Businesses", BIZ_COLS):
    prev = old.get(r["id"], {})
    b = {"id": r["id"], "name": r["name"], "host": r["host"], "sector": r["sector"],
         "e": r["emoji"] or c["sectors"].get(r["sector"], c["sectors"]["Other"])["emoji"],
         "col": r["color"] or "#FFC928", "place": r["place"],
         "dna": prev.get("dna") or c["sectors"].get(r["sector"], c["sectors"]["Other"])["dna"],
         "chips": [x.strip() for x in r["chips"].split(",") if x.strip()], "story": r["story"],
         "L": {k: r[k] for k in ("duration", "price", "availability", "activities") if r[k]},
         "cards": [], "qs": [], "bookings": prev.get("bookings", []), "threads": prev.get("threads", [])}
    if prev.get("why"):
        b["why"] = prev["why"]
    if r["map_x"] and r["map_y"]:
        b["pos"] = [float(r["map_x"]), float(r["map_y"])]
    biz.append(b)
by_id = {b["id"]: b for b in biz}
for r in rows("Postcards", CARD_COLS):
    if r["business_id"] not in by_id:
        sys.exit("Postcard from %s points at unknown business_id '%s'" % (r["name"], r["business_id"]))
    card = {"n": r["name"], "f": r["flag"], "l": r["language"], "t": r["text"], "bg": r["background"] or "sun", "s": r["stickers"]}
    if r["english"]: card["en"] = r["english"]
    if r["indonesian"]: card["idt"] = r["indonesian"]
    if truthy(r["local"]): card["loc"] = 1
    by_id[r["business_id"]]["cards"].append(card)
for r in rows("Questions", Q_COLS):
    if r["business_id"] not in by_id:
        sys.exit("Question points at unknown business_id '%s'" % r["business_id"])
    q = {"f": r["flag"], "t": r["text"]}
    if r["english"]: q["en"] = r["english"]
    if r["indonesian"]: q["idt"] = r["indonesian"]
    if truthy(r["local"]): q["loc"] = 1
    by_id[r["business_id"]]["qs"].append(q)
c["businesses"] = biz
c["phrases"] = [{"id": r["indonesian"], "en": r["english"], "say": r["pronunciation"], "emoji": r["emoji"]} for r in rows("Phrases", PHRASE_COLS)]
content_path.write_text(json.dumps(c, ensure_ascii=False, indent=1), encoding="utf-8")
print("Wrote data/content.json: %d businesses, %d postcards, %d questions, %d phrases" % (
    len(biz), sum(len(b["cards"]) for b in biz), sum(len(b["qs"]) for b in biz), len(c["phrases"])))
subprocess.run([sys.executable, str(ROOT / "scripts" / "export_static.py")], check=True)
sys.exit(subprocess.run([sys.executable, str(ROOT / "scripts" / "validate_data.py")]).returncode)
